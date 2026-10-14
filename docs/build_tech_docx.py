#!/usr/bin/env python3
"""Generate 快速且稳定车 technical Word doc from contest template structure."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent / "快速且稳定车_技术文档.docx"
OUT_DESK = Path.home() / "桌面" / "快速且稳定车_技术文档.docx"
IMG = Path(__file__).resolve().parent / "images"


def set_run_font(run, *, name="宋体", size=12, bold=False, color=None, code=False):
    run.bold = bold
    run.font.size = Pt(size)
    if code:
        run.font.name = "Consolas"
        r = run._element
        rPr = r.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:eastAsia"), "宋体")
        rFonts.set(qn("w:ascii"), "Consolas")
        rFonts.set(qn("w:hAnsi"), "Consolas")
    else:
        run.font.name = name
        r = run._element
        rPr = r.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:eastAsia"), name)
        rFonts.set(qn("w:ascii"), "Times New Roman")
        rFonts.set(qn("w:hAnsi"), "Times New Roman")
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, name="黑体", size=16 if level == 1 else (14 if level == 2 else 12), bold=True)
    return p


def add_para(doc, text, *, size=12, bold=False, first_indent=True, center=False):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    if first_indent and not center:
        pf.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_code(doc, code: str, *, title: str = ""):
    if title:
        p = doc.add_paragraph()
        run = p.add_run(f"【创新代码】{title}")
        set_run_font(run, name="黑体", size=11, bold=True, color=RGBColor(0xC0, 0x39, 0x2B))
    box = doc.add_paragraph()
    box.paragraph_format.line_spacing = 1.0
    box.paragraph_format.space_before = Pt(4)
    box.paragraph_format.space_after = Pt(8)
    box.paragraph_format.left_indent = Cm(0.3)
    # light highlight via shading
    shd = box._element.get_or_add_pPr()
    from docx.oxml import OxmlElement

    sh = OxmlElement("w:shd")
    sh.set(qn("w:fill"), "F5F5F5")
    sh.set(qn("w:val"), "clear")
    shd.append(sh)
    run = box.add_run(code.strip("\n") + "\n")
    set_run_font(run, size=9, code=True, color=RGBColor(0x1A, 0x1A, 0x2E))
    return box


def add_img(doc, name: str, caption: str, width_cm=14.5):
    path = IMG / name
    if not path.is_file():
        add_para(doc, f"（配图缺失：{name}）", first_indent=False)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_run_font(r, size=10, bold=True)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)

    # ===== Cover =====
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("全球校园人工智能算法精英大赛")
    set_run_font(r, name="黑体", size=18, bold=True)
    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t2.add_run("大模型技术创新赛 · 技术文档")
    set_run_font(r, name="黑体", size=16, bold=True)
    doc.add_paragraph()
    for line in [
        "作品名称：快速且稳定车",
        "学校名称：________________",
        "团队名称：快速且稳定车",
        "队长：________",
        "队员：________ / ________ / ________",
        "2026年10月",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        set_run_font(r, size=14)

    doc.add_page_break()

    # ===== Abstract =====
    add_heading(doc, "摘  要", 1)
    add_para(
        doc,
        "本作品「快速且稳定车」面向大模型技术创新赛取放任务，在 ROS2 Humble + Gazebo + Nav2 "
        "平台上构建「大模型解题—最短巡回规划—稳定导航—近场视觉伺服—区内投放」完整闭环。"
        "针对赛场动态障碍、字母投放区几何约束与多物块空驶成本三大难题，提出五项关键创新："
        "（1）DeepSeek 与本地正则求解器双通道交叉校验，并正确处理「都需要」人数倍增；"
        "（2）route_planner 联合优化物块编号绑定与访问顺序，最小化取放巡回长度；"
        "（3）Nav2 软成功（early_xy）+ MODE_NAV/MODE_LOCAL 权杖切换，避免物块旁爬行与空转；"
        "（4）按物块编号强制接近扇区与 standoff，保证前置相机 FOV 先看到目标；"
        "（5）投放区 zone_keepout 点云与已放块永久禁入，以及绿垫长边/短边非对称进区判定。"
        "整机不盲目抬高 max_vel，而以少空驶、少恢复行为、一次对准成功换取「快且稳」。",
    )
    add_para(
        doc,
        "关键词：ROS2；Nav2；OpenCV IBVS；最短巡回；区禁入；大模型解题；软成功切换",
        first_indent=False,
        bold=True,
    )

    doc.add_page_break()

    # ===== Ch1 =====
    add_heading(doc, "第一章 项目概述", 1)
    add_heading(doc, "1.1 项目背景与目标", 2)
    add_heading(doc, "1.1.1 背景介绍", 3)
    add_para(
        doc,
        "全球校园人工智能算法精英大赛「大模型技术创新赛」要求机器人根据自然语言应用题解析出 "
        "红/蓝物块数量与投放区（A/B/C），并在含动态障碍的室内地图中完成最多 5 块取放。"
        "仿真平台基于 Gazebo 与 Nav2，评分同时约束正确性、区内投放与总用时。作品名「快速且稳定车」"
        "体现设计原则：在定位可靠前提下压缩空驶与空等，同时保证抓取与投放一次成功。",
    )
    add_heading(doc, "1.1.2 核心问题", 3)
    add_para(
        doc,
        "核心难题包括：（1）题面表述多样，「都需要」易漏乘人数；（2）多块任务组合爆炸，随意顺序空驶长；"
        "（3）动态杆与膨胀层导致到点爬行、GoalAlign 空转；（4）特殊物块靠近墙/区时导航路径压块，相机看不到；"
        "（5）绿垫 1.0×0.5 m 长边/短边几何不同，垂直拧航向与进区平移易相互打架，块常落在短边外侧不得分。",
    )
    add_heading(doc, "1.1.3 项目目标", 3)
    add_para(doc, "目标1：题面→结构化任务 [颜色,数量,区] 准确率高，DeepSeek 失败可本地兜底。", first_indent=False)
    add_para(doc, "目标2：在 5 分钟量级内完成典型 5 块任务，空驶短、恢复行为少。", first_indent=False)
    add_para(doc, "目标3：抓取可靠、投放必须落入绿垫长方形内，且不碾压已放物块。", first_indent=False)

    add_heading(doc, "1.2 技术路线", 2)
    add_heading(doc, "1.2.1 任务分析", 3)
    add_para(
        doc,
        "系统任务按层级拆分：指令层解题与映射；任务层编号绑定与巡回排序；导航层 AMCL+Navfn+DWB；"
        "作业层 OpenCV HSV 近场对准 + LinkAttacher 吸附；投放层区内收与 keepout 反馈。大模型只输出符号计划，不进入控制环。",
    )
    add_heading(doc, "1.2.2 功能模块", 3)
    add_para(doc, "1. local_solver / llm_solver：自然语言 → (x,y) → /task_command", first_indent=False)
    add_para(doc, "2. route_planner：物块编号分配 + 最短访问顺序", first_indent=False)
    add_para(doc, "3. task_sm：状态机，调度取放腿与 MODE 切换", first_indent=False)
    add_para(doc, "4. motion：Nav2 封装、软成功、keepout、视觉 creep", first_indent=False)
    add_para(doc, "5. vision_align / vision_pose：HSV 检测与分区接近几何", first_indent=False)
    add_para(doc, "6. viz_bridge + coStudio：进度与状态面板", first_indent=False)

    add_heading(doc, "1.2.3 程序流程设计", 3)
    add_para(
        doc,
        "数据流：题面 → DeepSeek∥本地求解 → 任务表 → route_planner → task_sm → "
        "Nav2(AMCL+Navfn+DWB) → early_xy 软成功 → MODE_LOCAL → OpenCV IBVS → 吸附 → "
        "区内收投放 → zone_keepout/placed_keepout 更新 → 下一腿。",
    )
    add_img(doc, "code_pipeline.png", "图1 任务管线与状态机")

    # ===== Ch2 =====
    add_heading(doc, "第二章 功能设计", 1)
    add_heading(doc, "2.1 系统总体架构", 2)
    add_para(
        doc,
        "分层误差预算：全局用 known_coords 先验定位物块；中程用 AMCL 与代价地图；近场用像素偏差闭环；"
        "末端用距离门禁与吸附约束。速度方面保持 max_vel_x=1.35 m/s、max_vel_theta=0.80 rad/s，"
        "不把盲目提速作为主优化手段，避免定位未收敛时粒子发散。",
    )

    add_heading(doc, "2.2 关键技术详解（创新点）", 2)
    add_para(
        doc,
        "本节展开五项创新。对应实现代码见第三章「【创新代码】」高亮块，完整关键片段亦收录于第四章。",
        first_indent=False,
    )

    add_heading(doc, "2.2.1 双通道解题与「都需要」倍增", 3)
    add_para(
        doc,
        "官方规则：容量 5，x=ceil(类型1总量/5)，y=ceil(类型2总量/5)；数量≥3→高区 A，否则→低区 C。"
        "创新点在于：本地正则先处理「小帅、小丽都需要 N 件」按人数倍增，再剥离该子句避免重复累加；"
        "DeepSeek 与本地结果交叉校验，冲突时优先可执行的本地命令，保证赛场弱网仍可出题执行。",
    )
    add_img(doc, "score_mapping.png", "图2 题面到数量映射")

    add_heading(doc, "2.2.2 联合编号绑定与最短巡回", 3)
    add_para(
        doc,
        "route_planner 不只排序，还在可用立方块编号集合上枚举绑定，并对每组绑定求最短访问置换"
        "（N≤8 全排列，更大则最近邻）。代价函数计入 red#1→A/B 的门洞 via，从源头减少空驶与恢复。",
    )

    add_heading(doc, "2.2.3 软成功与 MODE 权杖", 3)
    add_para(
        doc,
        "取放接近段设置 early_xy：XY 进入邻域即 cancel NavigateToPose，切入 MODE_LOCAL 由视觉/开环完成厘米级对准。"
        "避免 GoalChecker 过严导致超时重规划，也避免 DWB 在高代价缝反复 GoalAlign 空转。作业结束再 exit_local 交回 Nav2。",
    )
    add_img(doc, "code_pick_flow.png", "图3 取物：软成功 → 视觉对准 → 吸附门禁")

    add_heading(doc, "2.2.4 分区强制接近几何", 3)
    add_para(
        doc,
        "创新点：取物目标不是物块坐标，而是保证 FOV 的 standoff 停靠点。对 red#2/#4、blue#4/#5 等危险编号"
        "强制西/东/北侧接近，禁止「为抄近路从对侧绕行压块」。yaw=atan2(块−车)，保证前置相机朝向目标。",
    )
    add_img(doc, "code_special_parks.png", "图4 特殊物块接近策略")
    add_img(doc, "code_known_cubes.png", "图5 物块坐标先验")

    add_heading(doc, "2.2.5 区禁入反馈与长/短边投放", 3)
    add_para(
        doc,
        "投放改变可通行空间：字母垫注入 zone_keepout 点云；当前区 set_keepout_except 允许进入，其他区保持禁入；"
        "已放块 add_placed_keepout 永久抬高代价。绿垫长边(X≈1.0m)/短边(Y≈0.5m) 非对称判定："
        "从 ±X 接近面对短边，必须沿长边越过短边进入垫内，避免落在短边外侧不得分；朝向大致对准即可，不与进区平移死磕垂直。",
    )

    # ===== Ch3 =====
    add_heading(doc, "第三章 功能实现", 1)
    add_heading(doc, "3.1 模块实现细节（创新代码精选）", 2)
    add_para(
        doc,
        "下列代码均来自本队工程实装，为评审突出「快且稳」相关创新点；完整扩展见第四章。",
        first_indent=False,
    )

    add_heading(doc, "3.1.1 本地求解器：都需要倍增 + ceil(/5)", 3)
    add_code(
        doc,
        '''def solve_xy_from_question(question: str) -> Tuple[int, int]:
    text = _normalize(question)
    type1 = type2 = 0
    _name = r"([\\u4e00-\\u9fffA-Za-z0-9、]{2,}?)"
    # 「小帅、小丽都需要7件T恤」→ 人数 × N
    for m in re.finditer(_name + r"都(?:需要|领养)(\\d+)件T恤", text):
        subjects = [s for s in m.group(1).split("、") if s]
        type1 += int(m.group(2)) * max(len(subjects), 1)
    # 剥离「都*」子句，避免下面再累加重复
    text2 = re.sub(_name + r"都(?:需要|领养)\\d+件T恤", "", text)
    for m in re.finditer(r"(\\d+)件T恤", text2):
        type1 += int(m.group(1))
    # …裤子/猫/狗同理 …
    x = int(math.ceil(type1 / 5.0))
    y = int(math.ceil(type2 / 5.0))
    return x, y


def map_xy_to_tasks(x, y, *, threshold=3, high_zone="A", low_zone="C"):
    def zone_for(n):
        return high_zone if n >= threshold else low_zone
    tasks = []
    if x > 0:
        tasks.append(PickTask("red", x, zone_for(x)))
    if y > 0:
        tasks.append(PickTask("blue", y, zone_for(y)))
    return tasks''',
        title="local_solver.py — 双通道可本地兜底的题面解析",
    )

    add_heading(doc, "3.1.2 最短巡回：门洞代价 + 编号绑定 + 置换搜索", 3)
    add_code(
        doc,
        '''def _leg_cost(pos, color, cid, zone, zones, *, use_door_via=True):
    """一段 pick→place：起点 →(门洞via)→ 物块 → 区心。"""
    cube = _cube_xy(color, cid)
    zxy = zones[zone.upper()]
    pts = [pos]
    if (use_door_via and color.lower() == "red"
            and int(cid) == 1 and zone.upper() in ("A", "B")):
        vias = list(_door_vias())  # (-4.2,0), (-3.2,0)
        dx, dy = cube[0] - pos[0], cube[1] - pos[1]
        L = math.hypot(dx, dy) or 1.0
        pts.extend(sorted(vias, key=lambda v:
            ((v[0]-pos[0])*dx + (v[1]-pos[1])*dy) / L))
    pts += [cube, zxy]
    return sum(_dist(a, b) for a, b in zip(pts, pts[1:])), zxy


def _best_order(legs, start, zones, use_door_via):
    """对已绑定编号的腿做最短访问顺序（N≤8 全排列）。"""
    best_total, best_legs = float("inf"), list(legs)
    if len(legs) <= 8:
        for perm in itertools.permutations(legs):
            total, _ = _tour_cost(perm, start, zones, use_door_via)
            if total < best_total:
                best_total, best_legs = total, list(perm)
    return best_total, best_legs''',
        title="route_planner.py — 空驶最短的联合规划",
    )

    add_heading(doc, "3.1.3 Nav2 软成功 early_xy", 3)
    add_code(
        doc,
        '''# motion.navigate：XY 进入邻域即取消 Nav2，交给近场，避免空转/爬行
while not result_fut.done():
    if time.time() - t0 > timeout:
        self.cancel_nav(); return False
    if early_xy_tol is not None:
        cur = self.get_base_xy_yaw_map()
        if cur is not None:
            dxy = math.hypot(cur[0] - float(x), cur[1] - float(y))
            if dxy <= float(early_xy_tol):
                self.cancel_nav()
                self._nav_goal_handle = None
                return True  # 软成功 → MODE_LOCAL
    time.sleep(0.05)


def enter_local_mode(self, settle_sec: float = 1.5) -> None:
    self.cancel_nav(); self.stop_chassis(n=2)
    time.sleep(min(0.06, max(float(settle_sec), 0.0)))
    self.drive_owner = "local"   # task_sm 拥有 cmd_vel

def exit_local_mode(self) -> None:
    self.stop_chassis(n=1)
    self.drive_owner = "nav"     # 交回 Nav2''',
        title="motion.py — early_xy 软成功与 MODE 权杖",
    )

    add_heading(doc, "3.1.4 特殊物块强制接近（保 FOV）", 3)
    add_code(
        doc,
        '''def approach_xy_for_cube(color, cube_id, *, standoff_m=0.35, robot_xy=None):
    cx, cy, _ = known_cube_xy(color, cube_id)
    so = float(standoff_m)
    def _face(ax, ay):
        return float(np.arctan2(cy - ay, cx - ax))

    if str(color).lower() == "blue":
        so_b = max(so, 0.60)
        if int(cube_id) == 5:          # 南走廊：北侧停靠，禁止压块
            ax, ay = float(cx), float(cy) + max(so_b, 0.72)
            return ax, ay, _face(ax, ay)
        if int(cube_id) == 4:          # 近 B 区：必须西侧，路径不穿块
            ax, ay = float(cx) - max(so_b, 0.70), float(cy)
            return ax, ay, _face(ax, ay)
        # 与车同侧停靠，绝不从对侧抄近路压块
        if robot_xy is not None and float(robot_xy[0]) >= float(cx):
            return float(cx + so_b), float(cy), _face(cx + so_b, cy)
        return float(cx - so_b), float(cy), _face(cx - so_b, cy)

    # red#2 禁东侧（落入 B 禁入）；red#4 禁西侧（贴墙 FOV 漂）
    so_r = max(float(standoff_m), 0.55)
    if int(cube_id) == 2:
        return float(cx) - so_r, float(cy), _face(cx - so_r, cy)
    if int(cube_id) == 4:
        return float(cx) + so_r, float(cy), _face(cx + so_r, cy)
    ...''',
        title="vision_pose.py — 分区 standoff，先看见再抓",
    )

    add_heading(doc, "3.1.5 HSV 饱和色门禁（拒墙面伪目标）", 3)
    add_code(
        doc,
        '''def _hsv_mask(hsv, color):
    if color.lower() == "red":
        m1 = cv2.inRange(hsv, (0, 145, 85), (10, 255, 255))
        m2 = cv2.inRange(hsv, (170, 145, 85), (180, 255, 255))
        return cv2.bitwise_or(m1, m2)   # 高饱和 → 塑料块，拒砖墙
    return cv2.inRange(hsv, (92, 85, 45), (132, 255, 255))


def _cube_like(area, bw, bh, cx, cy, w, h, mask, x, y, *, min_cy_frac=0.28):
    frac = area / float(max(w * h, 1))
    if frac > 0.06:          # 墙面色块过大
        return False
    if area < 8:             # 仍允许远距小块
        return False
    aspect = float(bw) / float(max(bh, 1))
    if aspect < 0.45 or aspect > 2.2:
        return False
    if cy < float(min_cy_frac) * h:   # 过高 → 非地板块
        return False
    return True''',
        title="vision_align.py — 低延迟近场感知，无 DNN",
    )

    add_heading(doc, "3.1.6 投放区点云禁入与已放块永久标记", 3)
    add_code(
        doc,
        '''def configure_zone_keepouts(self, zones, *, half_x=0.48, half_y=0.28, step=0.06):
    """字母垫矩形采样为 PointCloud2，注入代价地图。"""
    self._zone_keepout = {}
    for letter, xy in zones.items():
        cx, cy = float(xy[0]), float(xy[1])
        pts = [(x, y, 0.05)
               for x in frange(cx - half_x, cx + half_x, step)
               for y in frange(cy - half_y, cy + half_y, step)]
        self._zone_keepout[str(letter).upper()] = pts
    self._publish_zone_keepout()


def add_placed_keepout(self, x, y, *, radius=0.22, step=0.06):
    """已放块永久抬代价，避免下一腿碾压。"""
    pts = [p for p in disk_samples(x, y, radius, step)]
    self._placed_keepout.extend(pts)
    self._publish_zone_keepout()


def set_keepout_except(self, zone, *, clear_maps=True):
    """仅放开当前投放区，其它垫与已放块仍禁入。"""
    self._keepout_except = None if zone is None else zone.upper()
    self._publish_zone_keepout()''',
        title="motion.py — 投放几何对导航的闭环反馈",
    )

    add_heading(doc, "3.1.7 绿垫长边/短边非对称进区", 3)
    add_code(
        doc,
        '''def _place_drop_inward(self, tx, ty, zx, zy, *, zone=""):
    """长边 X≈1.0m 内收；短边 Y≈0.5m 更紧。"""
    inward = float(self.get_parameter("place_inward_m").value)  # 0.18
    max_dx, max_dy = 0.32, 0.12   # 长边 / 短边不同夹紧
    from_west = ...
    tx = tx + inward if from_west else tx - inward
    tx = 0.75 * tx + 0.25 * zx
    ty = 0.70 * ty + 0.30 * zy
    tx = zx + clip(tx - zx, -max_dx, max_dx)
    ty = zy + clip(ty - zy, -max_dy, max_dy)
    return tx, ty, yaw


def _drop_inside_pad(self, fx, fy, zx, zy, *, from_west=None):
    """必须越过短边、沿长边进入垫内，贴短边外侧不得分。"""
    if abs(fy - zy) > 0.14:          # 短边
        return False
    if abs(fx - zx) > 0.40:          # 长边
        return False
    if from_west is True and fx < zx - 0.18:
        return False                 # 仍在靠近侧短边外
    if from_west is False and fx > zx + 0.18:
        return False
    return True''',
        title="task_sm.py — 区内得分的几何门禁",
    )

    add_heading(doc, "3.2 系统集成与调试", 2)
    add_para(doc, "模块通信：/exam_question、/task_command、/task_status、/yzbot/zone_keepout、Nav2 action、cmd_vel。", first_indent=False)
    add_para(doc, "一键启动：cd ~/dev_ws/yzbot && ./run.sh（避开 conda，使用 /usr/bin/python3）。", first_indent=False)
    add_para(doc, "调试要点：RViz 2D Pose Estimate → coStudio 导入 costudio_yzbot.json → 6-Exam 空格出题。", first_indent=False)
    add_para(doc, "关键参数：max_vel_x=1.35，inflation_radius=0.55，grasp_standoff_m=0.55，place_inward_m=0.18，dyn_obs≈35s。", first_indent=False)

    # ===== Ch4 =====
    add_heading(doc, "第四章 创新代码汇编与优化说明", 1)
    add_para(
        doc,
        "本章按模板「其他内容」收录更完整的创新代码，便于与第三章对照。所有片段均摘自作品实装仓库，"
        "突出「快速（少空驶/少空等）」与「稳定（FOV/区内/禁碾压）」两条主线。",
    )

    add_heading(doc, "4.1 巡回规划主入口", 2)
    add_code(
        doc,
        '''def plan_route(command, *, start=(0.0, 0.0), max_cubes=5, use_door_via=True):
    tasks = parse_command(command, max_cubes=max_cubes)
    jobs = expand_jobs(tasks, max_cubes=max_cubes)
    zones = _load_zones()
    assignments = _materialize_assignments(jobs, max_cubes)
    best = None
    for legs in assignments:
        total, ordered, detail = _best_order(legs, start, zones, use_door_via)
        if best is None or total < best.total_m:
            best = PlanResult(legs=ordered, total_m=total, start=start, detail=detail)
    return best''',
        title="plan_route — 绑定×排序联合最优",
    )

    add_heading(doc, "4.2 导航参数一览（稳定优先）", 2)
    add_code(
        doc,
        '''Nav2:
  max_vel_x=1.35, max_vel_theta=0.80, controller_frequency=15Hz
  planner=NavfnPlanner, controller=DWBLocalPlanner
  inflation_radius=0.55, BaseObstacle.scale≈36
作业:
  grasp_standoff_m=0.55 / retry=0.62
  place_inward_m=0.18, placed_keepout_radius≈0.22
  MODE_NAV / MODE_LOCAL 权杖, early_xy 软成功
动态障碍: 正弦周期约 35s（原速），以代价地图+DWB 绕行''',
        title="不拿 max_vel 换虚快，用结构换真快",
    )

    add_heading(doc, "4.3 创新点小结", 2)
    add_para(doc, "1. 双通道解题 +「都需要」倍增，弱网可跑。", first_indent=False)
    add_para(doc, "2. 编号绑定与最短巡回，显著砍空驶。", first_indent=False)
    add_para(doc, "3. early_xy + MODE 切换，去爬行/空转。", first_indent=False)
    add_para(doc, "4. 特殊块强制接近几何，先看见再抓。", first_indent=False)
    add_para(doc, "5. keepout 闭环 + 长/短边区内门禁，投放可得分且不碾块。", first_indent=False)

    add_heading(doc, "参考文献", 1)
    refs = [
        "[1] ROS 2 Documentation — Humble. https://docs.ros.org/en/humble/",
        "[2] Nav2 — Navigation Framework for ROS 2. https://navigation.ros.org/",
        "[3] Fox D. et al. Monte Carlo Localization. AAAI/IROS related works on AMCL.",
        "[4] Fox D., Burgard W., Thrun S. The Dynamic Window Approach to Collision Avoidance.",
        "[5] OpenCV Tutorials — Image Processing (HSV thresholding, contours).",
        "[6] 全球校园人工智能算法精英大赛 · 大模型技术创新赛 比赛指导资料（2026）.",
    ]
    for r in refs:
        add_para(doc, r, first_indent=False)

    doc.save(OUT)
    try:
        doc.save(OUT_DESK)
        desk = str(OUT_DESK)
    except Exception as exc:  # noqa: BLE001
        desk = f"(桌面保存失败: {exc})"
    print("WROTE", OUT)
    print("DESK", desk)


if __name__ == "__main__":
    build()
